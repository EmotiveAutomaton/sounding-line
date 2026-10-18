"""Source-only Jester Dataset2+ intake; no liking/model comparison is computed.

DESIGN CHECK: LESSONS sections3-5 read October18 for this intake. Known valid
source row counts, numeric cell types, missing sentinel, column identity and
exact ratings pass the source gate. Changed, malformed or missing required
fields refuse intake; none is a positive, null or negative scientific verdict.
No failed/missing row is discarded.
The official archive's DIMENSIONS contain 50,692 rows, whereas its archived
page advertises 66,336. Both facts remain in the report; newer Dataset3 ratings
are not substituted. All 150 joke texts/columns survive intake.
"""
from __future__ import annotations

import importlib
import io
import math
from pathlib import Path
import sys
import time
import zipfile

import numpy as np

from .timing_census import filehash, read, write
from .timing_p6 import progress, write_equal
from .timing_stats import bind, check_binding

CONFIG = dict(card="Jester2plus-intake", rows_per_chunk=1024)
INPUTS = frozenset(("ratings_archive", "source_page", "parser_wheel",
                    "texts_archive", "jester_texts", "source_manifest"))
FORECAST_SECONDS = 300
CUMULATIVE_CAP_SECONDS = 900
WHEEL_SHA = "ea762c3d29f4cca48d82df517b6d89fbce4db3107f9d78713e48cd321d5c9aa9"
ARCHIVE_SHA = "de360773d91253200a203afa99c38d3368bfcfcc02d28088f11d5ec1dacc9f6f"
WORKBOOK_SHA = "41db5a229974f41d6eca9f392913db485a242ff93393f46b03f1fe56d0480b4b"
TEXT_ARCHIVE_SHA = "e4afaada9bd3b04cd614b383812ec5f911f758198d0f175f8957f2b42f111047"
TEXTS_SHA = "e7d00d7091ffc826619f25e634d89e1e07b2dbe8aee126474dc201b61a58f94a"
DOCUMENTED_UNRATED = [1, 2, 3, 4, 6, 9, 10, 11, 12, 14]


def sha_bytes(value):
    import hashlib
    return hashlib.sha256(value).hexdigest()


def validate(inputs, config):
    if config != CONFIG or type(config.get("rows_per_chunk")) is not int:
        raise ValueError("exact Jester intake configuration required")
    if set(inputs) != INPUTS:
        raise ValueError("exact Jester intake inputs required")
    m = read(inputs["source_manifest"])
    expected = {k: filehash(inputs[k]) for k in sorted(INPUTS - {"source_manifest"})}
    if m.get("schema") != 1 or m.get("input_sha256") != expected:
        raise ValueError("Jester source manifest/input identity differs")
    if expected["parser_wheel"] != WHEEL_SHA:
        raise ValueError("isolated xlrd dependency differs")
    if m.get("source_kind") not in ("official_dataset2plus", "constructed_fixture"):
        raise ValueError("unrecognized Jester source kind")
    if m.get("source_kind") == "official_dataset2plus":
        if (expected["ratings_archive"] != ARCHIVE_SHA or m.get("workbook_sha256") != WORKBOOK_SHA
                or expected["texts_archive"] != TEXT_ARCHIVE_SHA or expected["jester_texts"] != TEXTS_SHA
                or m.get("ratings_member") != "jesterfinal151cols.xls"
                or m.get("sheets") != [dict(name="jesterfinal151cols", rows=50692, columns=151)]
                or m.get("advertised_raters") != 66336):
            raise ValueError("official Dataset2+ source replaced or population changed")
    sheets = m.get("sheets", [])
    if (not sheets or len({s["name"] for s in sheets}) != len(sheets)
            or any(type(s["rows"]) is not int or s["rows"] <= 0 or s["columns"] != 151 for s in sheets)):
        raise ValueError("invalid declared worksheet population")
    return m, dict(config=config, inputs={k: filehash(inputs[k]) for k in sorted(inputs)})


def parser(wheel):
    """Zipimport the pinned private wheel; never install into the live environment."""
    wheel = str(Path(wheel).resolve())
    if filehash(wheel) != WHEEL_SHA:
        raise ValueError("parser wheel changed")
    if wheel not in sys.path:
        sys.path.insert(0, wheel)
    xlrd = importlib.import_module("xlrd")
    if xlrd.__version__ != "2.0.2" or not str(xlrd.__file__).replace("\\", "/").startswith(wheel.replace("\\", "/") + "/"):
        raise ValueError("xlrd did not load from the pinned isolated wheel")
    return xlrd


def member(archive, name):
    with zipfile.ZipFile(archive) as z:
        entries = [i for i in z.infolist() if i.filename == name]
        if len(entries) != 1 or entries[0].file_size > 64 * 1024**2:
            raise ValueError("missing, repeated or oversized archive member")
        return z.read(entries[0])


def validate_texts(inputs, manifest):
    """Independently reconstruct the exact ID-to-text mapping from its workbook."""
    import openpyxl
    raw = member(inputs["texts_archive"], manifest["texts_member"])
    if sha_bytes(raw) != manifest["texts_workbook_sha256"]:
        raise ValueError("text workbook identity differs")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        if wb.sheetnames != [manifest["texts_sheet"]]:
            raise ValueError("text worksheet population differs")
        sheet = wb[manifest["texts_sheet"]]
        values = []
        trailing_blank = 0
        for index, row in enumerate(sheet.iter_rows(), 1):
            if index <= 150:
                if not row or row[0].data_type == "f" or not isinstance(row[0].value, str) or not row[0].value.strip():
                    raise ValueError("missing, formula or non-text joke cell")
                values.append(dict(joke_id=index, text=row[0].value))
                if any(cell.value is not None for cell in row[1:]):
                    raise ValueError("unexpected text workbook columns")
            elif any(cell.value is not None for cell in row):
                raise ValueError("unexpected populated joke row beyond 150")
            else:
                trailing_blank += 1
        if len(values) != 150 or read(inputs["jester_texts"]) != values:
            raise ValueError("joke text/ID differs from exact A1:A150 source cells")
        return dict(jokes=150, ids=list(range(1, 151)), trailing_blank_rows=trailing_blank,
                    texts_sha256=filehash(inputs["jester_texts"]), normalization="none")
    finally:
        wb.close()


def open_source(inputs, manifest):
    raw = member(inputs["ratings_archive"], manifest["ratings_member"])
    if sha_bytes(raw) != manifest["workbook_sha256"]:
        raise ValueError("rating workbook identity differs")
    log = io.StringIO()
    book = parser(inputs["parser_wheel"]).open_workbook(file_contents=raw, on_demand=True,
        logfile=log, formatting_info=False, ignore_workbook_corruption=False)
    if book.biff_version != 80 or book.sheet_names() != [s["name"] for s in manifest["sheets"]]:
        book.release_resources()
        raise ValueError("rating workbook BIFF version or worksheets differ")
    return book, log


def validate_rows(values, types):
    """Only numeric [-10,10] or exact99; numeric count must equal nonmissing cells."""
    a = np.asarray(values, dtype=np.float64)
    t = np.asarray(types)
    if a.ndim != 2 or a.shape[1] != 151 or a.shape != t.shape or not np.all(t == 2):
        raise ValueError("rating worksheet requires exactly151 numeric cells per row")
    if not np.isfinite(a).all():
        raise ValueError("nonfinite rating source cell")
    ratings = a[:, 1:]
    observed = ratings != 99.
    if np.any(observed & ((ratings < -10.) | (ratings > 10.))):
        raise ValueError("rating outside [-10,10] and missing sentinel99")
    counts = observed.sum(1)
    if not np.array_equal(a[:, 0], counts):
        raise ValueError("first-column count does not equal observed ratings")
    result = np.array(ratings, dtype=np.float64, order="C", copy=True)
    result[~observed] = np.nan
    return result


def source_chunks(book, manifest):
    offset = 0
    for number, spec in enumerate(manifest["sheets"]):
        sheet = book.sheet_by_index(number)
        if sheet.nrows != spec["rows"] or sheet.ncols != 151:
            raise ValueError("actual worksheet dimensions differ from frozen metadata")
        for start in range(0, sheet.nrows, CONFIG["rows_per_chunk"]):
            stop = min(start + CONFIG["rows_per_chunk"], sheet.nrows)
            array = validate_rows([sheet.row_values(i) for i in range(start, stop)],
                                  [sheet.row_types(i) for i in range(start, stop)])
            yield dict(sheet=spec["name"], sheet_index=number, excel_row_first=start + 1,
                       excel_row_last=stop, rater_offset=offset + start, rows=stop - start), array
        offset += sheet.nrows
        book.unload_sheet(number)


def statistics(manifest, texts, count, observed_counts, zero_raters, parser_messages):
    unrated = [i + 1 for i, n in enumerate(observed_counts) if n == 0]
    return dict(card=CONFIG["card"], state="positive", instrument_only=True,
        scientific_admission=False, source_kind=manifest["source_kind"], raters=count,
        advertised_raters=manifest["advertised_raters"],
        advertised_count_matches=(count == manifest["advertised_raters"]),
        population="Every numeric source row retained; row identity is worksheet plus original Excel row. No rater deduplication or outcome filter.",
        rating_values="Original float64 numeric cells preserved; exact99 becomes missing NaN, never zero.",
        first_column_count_verified=True, zero_rating_raters=zero_raters,
        texts=texts, joke_observed_rating_counts=[int(n) for n in observed_counts],
        jokes_with_ratings=150 - len(unrated), unrated_joke_ids=unrated,
        documented_unrated_joke_ids=DOCUMENTED_UNRATED,
        documented_unrated_matches=(unrated == DOCUMENTED_UNRATED),
        eligibility="All150 texts retained for P7; only columns with observed ratings can enter D36 liking contrasts. Unrated is not a zero rating.",
        parser_messages=parser_messages,
        source_description_discrepancy="Archived page advertises66336 rows; pinned official workbook metadata declares50692. Preserve actual full workbook population; no Dataset3 substitution." if manifest["source_kind"] == "official_dataset2plus" else "Constructed known-answer fixture.")


def execute(output: Path, inputs: dict, config: dict) -> dict:
    started = time.monotonic()
    output = Path(output)
    manifest, binding = validate(inputs, config)
    if (output / "TIMING_BINDINGS.json").exists():
        return verify(output, inputs, config)["primary"]
    write_equal(output / "INPUT_BINDING.json", binding)
    old = read(output / "PROGRESS.json") if (output / "PROGRESS.json").exists() else {}
    prior = old.get("active_seconds", 0.)
    if not isinstance(prior, (int, float)) or not math.isfinite(prior) or prior < 0:
        raise ValueError("invalid intake cumulative time")
    elapsed = lambda: prior + time.monotonic() - started
    texts = validate_texts(inputs, manifest)
    book, parser_log = open_source(inputs, manifest)
    total = sum(s["rows"] for s in manifest["sheets"])
    counts = np.zeros(150, dtype=np.int64)
    row_count = zero_raters = 0
    index = []
    try:
        for unit, array in source_chunks(book, manifest):
            name = "chunks/" + str(len(index)).zfill(4) + ".npy"
            dest = output / name
            receipt_path = dest.with_suffix(".json")
            if receipt_path.exists():
                saved = read(receipt_path)
                if saved != dict(unit=unit, input_binding_sha256=filehash(output / "INPUT_BINDING.json"), sha256=filehash(dest)):
                    raise ValueError("saved intake source or chunk identity differs")
                if not np.array_equal(np.load(dest, allow_pickle=False), array, equal_nan=True):
                    raise ValueError("saved ratings differ from reconstructed source")
            else:
                if dest.exists():
                    raise ValueError("unsealed ratings chunk exists")
                dest.parent.mkdir(parents=True, exist_ok=True)
                with dest.open("xb") as f:
                    np.save(f, array, allow_pickle=False)
                write(receipt_path, dict(unit=unit, input_binding_sha256=filehash(output / "INPUT_BINDING.json"), sha256=filehash(dest)))
            seen = np.isfinite(array)
            counts += seen.sum(0)
            zero_raters += int(np.sum(~seen.any(1)))
            row_count += len(array)
            index.append(dict(file=name, receipt=receipt_path.relative_to(output).as_posix(), **unit))
            progress(output / "PROGRESS.json", dict(phase="parse", complete_raters=row_count,
                total_raters=total, active_seconds=elapsed()))
            if row_count >= math.ceil(total / 10) and not (output / "FIRST_TENTH.json").exists():
                write(output / "FIRST_TENTH.json", dict(complete_raters=row_count, total_raters=total,
                    active_seconds=elapsed(), forecast_cpu_seconds=FORECAST_SECONDS,
                    projected_cpu_seconds=2 * elapsed() * total / row_count,
                    method="Conservative full elapsed per completed row, including workbook load and text/source validation, doubled for fresh source verification.",
                    output_bytes_so_far=sum(p.stat().st_size for p in output.rglob("*") if p.is_file())))
            if elapsed() >= CUMULATIVE_CAP_SECONDS:
                raise ValueError("Jester cumulative source-intake cap reached")
            if (output / "YIELD_REQUEST.json").exists():
                return dict(status="yielded", phase="parse", complete_raters=row_count, total_raters=total)
    finally:
        book.release_resources()
    if row_count != total:
        raise ValueError("not every source rater was retained")
    result = statistics(manifest, texts, row_count, counts, zero_raters, parser_log.getvalue())
    write_equal(output / "RATINGS_INDEX.json", dict(schema=1, joke_ids=list(range(1, 151)),
        missing="NaN", dtype="float64", rows=total, chunks=index))
    write_equal(output / "PRIMARY.json", result)
    artifacts = ["INPUT_BINDING.json", "RATINGS_INDEX.json", "PRIMARY.json", "FIRST_TENTH.json"]
    artifacts += [r[k] for r in index for k in ("file", "receipt")]
    bind(output, inputs, config, artifacts)
    progress(output / "PROGRESS.json", dict(phase="complete", complete_raters=row_count,
        total_raters=total, active_seconds=elapsed()))
    return result


def verify(output: Path, inputs: dict, config: dict) -> dict:
    """Fresh-process reparse of source cells, not reconstruction from saved ratings."""
    started = time.monotonic()
    output = Path(output)
    manifest, binding = validate(inputs, config)
    check_binding(output, inputs, config)
    if read(output / "INPUT_BINDING.json") != binding:
        raise ValueError("input binding changed")
    prior = read(output / "PROGRESS.json")["active_seconds"]
    if not math.isfinite(prior) or prior < 0:
        raise ValueError("invalid intake cumulative time")
    texts = validate_texts(inputs, manifest)
    book, parser_log = open_source(inputs, manifest)
    counts = np.zeros(150, dtype=np.int64)
    row_count = zero_raters = 0
    index = []
    try:
        for unit, array in source_chunks(book, manifest):
            name = "chunks/" + str(len(index)).zfill(4) + ".npy"
            dest = output / name
            receipt_path = dest.with_suffix(".json")
            expected = dict(unit=unit, input_binding_sha256=filehash(output / "INPUT_BINDING.json"), sha256=filehash(dest))
            if read(receipt_path) != expected or not np.array_equal(np.load(dest, allow_pickle=False), array, equal_nan=True):
                raise ValueError("saved ratings or identity differ from fresh source")
            seen = np.isfinite(array)
            counts += seen.sum(0)
            zero_raters += int(np.sum(~seen.any(1)))
            row_count += len(array)
            index.append(dict(file=name, receipt=receipt_path.relative_to(output).as_posix(), **unit))
            if prior + time.monotonic() - started >= CUMULATIVE_CAP_SECONDS:
                raise ValueError("Jester cumulative cap reached during fresh verification")
    finally:
        book.release_resources()
    expected_index = dict(schema=1, joke_ids=list(range(1, 151)), missing="NaN", dtype="float64", rows=row_count, chunks=index)
    primary = statistics(manifest, texts, row_count, counts, zero_raters, parser_log.getvalue())
    if read(output / "RATINGS_INDEX.json") != expected_index or read(output / "PRIMARY.json") != primary:
        raise ValueError("Jester report differs from fresh complete source reconstruction")
    return dict(verified=True, primary=primary, source_raters_reparsed=row_count,
                cumulative_seconds=prior + time.monotonic() - started)
